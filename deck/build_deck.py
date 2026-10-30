"""Fill the challenge's own six-slide template in place, so the deck keeps its
layout and colours. Text is set in Arial: the template asks for Nirmala UI,
which is not on every machine that will render the PDF, and a substituted font
with different widths would overflow the template's small boxes."""
import copy, os, subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "..", "figures")
OUT = os.path.join(HERE, "VIGIL-HEISENBUG-Seva2026")
DEMO = "amayakavya.github.io/vigil"

NAVY, ORANGE, INK, GREY, PALE, WHITE, PEACH = (RGBColor.from_string(c) for c in
    ("1B2A5C", "F28C28", "1F2937", "6B7280", "C9D2EC", "FFFFFF", "FDEBD7"))
FONT = "Arial"

prs = Presentation(os.path.join(HERE, "template.pptx"))
S = prs.slides


def sh(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def drop(slide, *names):
    for n in names:
        e = sh(slide, n)._element
        e.getparent().remove(e)


def put(shape, paras, size=9.5, color=INK, bold=False, italic=False, align=None, space=2, anchor=None):
    """paras: a string (lines split on newline) or a list of paragraphs, each a
    string or a list of (text, bold) runs."""
    tf = shape.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = paras.split("\n")
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        if first:
            # The template's guidance paragraphs carry bullets and hanging
            # indents. Start from a bare paragraph so every line matches.
            for child in list(p._p):
                p._p.remove(child)
        first = False
        if align:
            p.alignment = align
        p.space_after = Pt(space)
        runs = para if isinstance(para, list) else [(para, bold)]
        for text, b in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name, f.size, f.bold, f.italic = FONT, Pt(size), b, italic
            f.color.rgb = color


def box(slide, x, y, w, h, paras, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.margin_left = tb.text_frame.margin_right = Inches(0.02)
    tb.text_frame.margin_top = tb.text_frame.margin_bottom = Inches(0.0)
    put(tb, paras, **kw)
    return tb


def card(slide, x, y, w, h, fill=WHITE, radius=0.06):
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.adjustments[0] = radius
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    r.line.fill.background()
    r.shadow.inherit = False
    return r


def tick(shape, mark="✓"):
    shape.fill.solid()
    shape.fill.fore_color.rgb = ORANGE
    shape.line.color.rgb = ORANGE
    put(shape, mark, size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    shape.text_frame.margin_left = shape.text_frame.margin_right = 0
    shape.text_frame.margin_top = shape.text_frame.margin_bottom = 0


def sources(slide, text):
    box(slide, 0.5, 5.2, 5.4, 0.3, text, size=6.5, color=GREY)


# 1 · Title and snapshot ---------------------------------------------------
s = S[0]
put(sh(s, "Text 0"), "SEVA FIRST INNOVATION CHALLENGE 2026  •  DEFENCE  •  PREDICTIVE MAINTENANCE, MRO AND LIFECYCLE INTELLIGENCE",
    size=9, color=ORANGE, bold=True)
put(sh(s, "Text 1"), "VIGIL", size=38, color=WHITE, bold=True)
box(s, 0.5, 1.5, 5.6, 0.3, "VIBRATION INTELLIGENCE GRID FOR INTEGRITY AND LIFECYCLE", size=9.5, color=PALE, bold=True)
put(sh(s, "Text 2"), [[("Pulse diagnosis for machines. ", True)],
                      [("A clip-on vibration sensor that tells a police or forces mechanic which part of a generator, "
                        "pump or vehicle is failing, days before it fails. It needs no internet.", False)]],
    size=11, color=PALE, space=3)
sh(s, "Text 2").height = Inches(0.75)
put(sh(s, "Text 16"), "Measured so far", size=13, color=ORANGE, bold=True)
t17 = sh(s, "Text 17")
t17.height = Inches(2.7)
put(t17, [
    [("4 of 4 ", True), ("bearing failures warned, 42 to 603 h before the test rig was stopped", False)],
    [("4 of 4 ", True), ("failing parts named correctly", False)],
    [("0 ", True), ("alarms not traced to the failing bearing in 5,366 healthy bearing-hours", False)],
    [("40 of 52 ", True), ("seeded faults named on a second dataset. 2 wrong, 10 declined", False)],
    [("NASA IMS and Case Western public bearing data. Hardware node not yet built.", False)],
], size=9.5, color=WHITE, space=6)
t17.text_frame.paragraphs[-1].runs[0].font.color.rgb = PALE
t17.text_frame.paragraphs[-1].runs[0].font.size = Pt(8)
card(s, 0.5, 2.85, 5.65, 2.42)
s.shapes.add_picture(os.path.join(FIG, "ims_run2.png"), Inches(0.58), Inches(2.9), width=Inches(5.5))
box(s, 6.4, 4.4, 2.15, 0.95, [[("Team HEISENBUG", True)], [("Amaya Kavya, team lead", False)],
                              [("EXTC, RAIT, Navi Mumbai", False)], [("Live demo", True)], [(DEMO, False)]],
    size=8, color=PALE, space=1)
s.shapes.add_picture(os.path.join(FIG, "qr.png"), Inches(8.62), Inches(4.38), width=Inches(0.88))

# 2 · Problem ----------------------------------------------------------------
s = S[1]
put(sh(s, "Text 7"), "Generators, pumps and vehicles are serviced by the calendar. Healthy parts get replaced on schedule, "
    "and a part that starts failing between services is not seen until the machine stops. Seeing inside a running "
    "machine today takes a costly analyser and a trained analyst, and the workshops have neither.", size=8.5)
put(sh(s, "Text 10"), "Motor Transport sections of Maharashtra Police (Home Department), which keep the force's vehicles "
    "and station generators running, and the personnel who depend on them. The same gap sits in forces MRO units. "
    "Proposed pilot: one MT section, Navi Mumbai.", size=8.5)
put(sh(s, "Text 13"), [
    "Maharashtra Police run 15,823 vehicles, including 3,281 special vehicles such as 72 mine-protected ones (BPR&D, 1 Jan 2024).",
    "Bearings cause 41 to 44% of motor failures (EPRI 1982, IEEE 1985).",
    "The Indian Navy has put AI-based predictive maintenance out as an iDEX challenge (DISC XI)."], size=8.2, space=3)
put(sh(s, "Text 16"), [
    "Predictive maintenance cuts breakdowns by 70 to 75% (US DOE, 2010). The same guide notes its equipment often costs over $50,000 and needs trained staff.",
    "Commercial systems mostly report to a vendor's cloud, which a security force cannot use."], size=8.2, space=3)
put(sh(s, "Text 18"), "15,823", size=30, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
put(sh(s, "Text 19"), "vehicles with Maharashtra Police, plus the generators and pumps behind its 1,193 police stations",
    size=9, color=WHITE, align=PP_ALIGN.CENTER)
sh(s, "Text 19").height = Inches(0.85)
put(sh(s, "Text 21"), "Each breakdown takes a vehicle off the road, or a station off backup power, until the workshop clears it.", size=9,
    color=PALE, align=PP_ALIGN.CENTER)
sources(s, "Sources: BPR&D, Data on Police Organisations as on 01.01.2024, Tables 5.1.6, 5.1.6(A), 5.1.9. EPRI EL-2678 (1982). "
           "IEEE-IAS Motor Reliability Working Group (1985). US DOE FEMP, O&M Best Practices Guide R3.0 (2010), sec. 5.4. iDEX DISC XI results.")

# 3 · Solution and innovation ------------------------------------------------
s = S[2]
put(sh(s, "Text 6"), "A clip-on node with a vibration sensor and its own processor. It learns each machine's normal in its "
    "first day, then checks every reading against it and against the defect frequencies of that machine's bearings. "
    "When a defect line rises, it names the part and the spare to indent. Results reach a workshop console over "
    "Bluetooth. Nothing goes over the internet.", size=8.5)
sh(s, "Text 6").height = Inches(0.95)
put(sh(s, "Text 12"), "1 s of vibration every 10 min, and the bearing number from the asset register", size=7.5)
put(sh(s, "Text 18"), "Compare with the machine's own normal, then check the bearing's defect lines", size=7.5)
put(sh(s, "Text 24"), "“Alternator bearing 6309, outer race: replace at next stop”", size=7.5)
tbl = next(x for x in s.shapes if x.has_table).table
rows = [
    ("Analyser often over $50,000, plus a trained analyst", "About ₹4,000 a point (estimate), no subscription"),
    ("Found at the next service, or at breakdown", "Warned 42 to 603 h ahead on NASA failure data"),
    ("Vendor cloud, or a visiting expert", "Fully offline, inside the premises"),
    ("Spectra that an analyst must read", "Names the part and the spare in plain words"),
]
for i, (a, b) in enumerate(rows, start=1):
    for j, text in ((1, a), (2, b)):
        cell = tbl.cell(i, j)
        put(cell, text, size=7.5, color=INK, bold=(j == 2))
put(sh(s, "Text 27"), "It names the failing part, working from the bearing's own geometry, so it needs no failure history "
    "to learn from. It also runs without the internet. Anyone can rerun every number on the demo page from public data.", size=8.5)
sh(s, "Text 27").height = Inches(0.8)
sources(s, "Existing-approach cost: US DOE FEMP, O&M Best Practices Guide R3.0 (2010), sec. 5.4. Demo and code: " + DEMO)

# 4 · Feasibility -------------------------------------------------------------
s = S[3]
tick(sh(s, "Shape 7"))
tick(sh(s, "Shape 9"))
put(sh(s, "Text 10"), "Proof of Concept", size=9, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
box(s, 1.6, 2.47, 7.0, 0.22, "Method proven on public laboratory data. Hardware node designed, not yet built: bench build from November 2026.",
    size=7.5, color=GREY, italic=True, align=PP_ALIGN.CENTER)
put(sh(s, "Text 19"), [
    "Node: ESP32-S3 and ST IIS3DWB accelerometer (flat to 6 kHz), sealed housing, magnet or stud mount, 12/24 V.",
    "Gateway: small Linux board with Bluetooth.",
    "Analysis: Python, ported to C on the node.",
    "Bench: seeded-fault rig at the college EXTC lab."], size=7.4, space=2)
put(sh(s, "Text 22"), [
    "Node about ₹4,000, gateway about ₹15,000 (retail estimates).",
    "Pilot, 60 points, 6 months: ₹10 lakh.",
    "Hardware 3.7 L. Two field technicians 2.4 L. Test rig, IP and EMC checks 1.5 L. Software and security review 1.2 L. Travel 0.5 L. Contingency 0.7 L."],
    size=7.4, space=2)
put(sh(s, "Text 25"), [
    [("3 months: ", True), ("10 nodes bench tested, asset survey done, 40 points on generators and pumps.", False)],
    [("6 months: ", True), ("20 points on vehicles at idle. Warnings checked against strip-down findings.", False)],
    [("12 months: ", True), ("report to the Home Department, second district, iDEX entry for forces MRO.", False)]],
    size=7.4, space=3)
put(sh(s, "Text 28"), [
    "Pilot host (proposed): Maharashtra Police Motor Transport, Home Department.",
    "Facilitation: MSInS.",
    "Scale-up: SRPF, then Army, Navy and Air Force MRO units through iDEX."], size=7.4, space=3)

# 5 · Impact and scale ---------------------------------------------------------
s = S[4]
put(sh(s, "Text 6"), "60", size=30, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
put(sh(s, "Text 7"), "Machines monitored", size=11, color=INK, bold=True, align=PP_ALIGN.CENTER)
put(sh(s, "Text 8"), "in one MT section in year one, and the crews and stations they serve", size=7.5, color=GREY,
    italic=True, align=PP_ALIGN.CENTER)
sh(s, "Text 8").height = Inches(0.4)
put(sh(s, "Text 10"), "25 to 30%", size=30, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
put(sh(s, "Text 11"), "Maintenance cost cut", size=11, color=INK, bold=True, align=PP_ALIGN.CENTER)
put(sh(s, "Text 12"), "reported for predictive programmes (US DOE, 2010). The pilot measures ours.", size=7.5, color=GREY,
    italic=True, align=PP_ALIGN.CENTER)
sh(s, "Text 12").height = Inches(0.4)
put(sh(s, "Text 14"), "2 + 10", size=30, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
put(sh(s, "Text 15"), "Jobs and skills", size=11, color=INK, bold=True, align=PP_ALIGN.CENTER)
put(sh(s, "Text 16"), "technician jobs, and MT mechanics trained in condition monitoring", size=7.5, color=GREY,
    italic=True, align=PP_ALIGN.CENTER)
sh(s, "Text 16").height = Inches(0.4)
put(sh(s, "Text 19"), [
    "A police vehicle that does not break down on duty reaches the 112 call it was sent to.",
    "Fewer breakdowns also leave fewer personnel stranded in the field. Parts get changed when they are actually worn, so less goes to scrap."], size=8, space=3)
put(sh(s, "Text 22"), [
    "Paid from the maintenance budget the department already spends, as a yearly fee per point, priced below what one avoided breakdown saves.",
    "No cloud subscription. Nodes built and serviced in Maharashtra."], size=8, space=3)
put(sh(s, "Text 25"), [
    "One MT section first, then every district MT section and the SRPF.",
    "The forces' MRO units come through iDEX, where the Navy has already asked for this.",
    "The same node goes everywhere. Each site needs only its own asset register."], size=8, space=3)

# 6 · Team and support ----------------------------------------------------------
s = S[5]
drop(s, *[f"Shape {i}" for i in (10, 11, 16, 17, 22, 23)], *[f"Text {i}" for i in (12, 13, 14, 15, 18, 19, 20, 21, 24, 25, 26, 27, 28)])
card_shape = sh(s, "Shape 4")
card_shape.width = Inches(4.75)
put(sh(s, "Text 6"), "AK", size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
put(sh(s, "Text 7"), "Amaya Kavya", size=11.5, color=WHITE, bold=True)
sh(s, "Text 7").width = Inches(3.6)
put(sh(s, "Text 8"), "Team lead, solo entry", size=8.5, color=ORANGE, italic=True)
sh(s, "Text 8").width = Inches(3.6)
put(sh(s, "Text 9"), "B.E. Electronics and Telecommunication, second year, Ramrao Adik Institute of Technology, Navi Mumbai. "
    "Signal processing and embedded systems. Intern, SBI Global IT Centre, 2026.", size=8, color=PALE)
sh(s, "Text 9").width = Inches(3.75)
sh(s, "Text 9").height = Inches(0.55)
box(s, 0.5, 2.72, 4.8, 0.25, "Working demo, open code", size=10.5, color=ORANGE, bold=True)
s.shapes.add_picture(os.path.join(FIG, "replay-shot.png"), Inches(0.5), Inches(3.02), width=Inches(2.95))
s.shapes.add_picture(os.path.join(FIG, "qr.png"), Inches(3.8), Inches(3.02), width=Inches(1.0))
box(s, 3.8, 4.07, 1.5, 0.55, [DEMO.split("/")[0] + "/", DEMO.split("/")[1]], size=7.5, color=PALE, space=0)
for shape_name in ("Shape 31", "Shape 33", "Shape 35", "Shape 39"):
    tick(sh(s, shape_name))
put(sh(s, "Text 41"), "The method works on laboratory data. A pilot with one Police Motor Transport section is how we find "
    "out what it does on Maharashtra's own machines.", size=8.5, color=INK, italic=True)
sh(s, "Text 41").height = Inches(0.75)
put(sh(s, "Text 42"), "Thank you  |  kav.ama.rt25@dypatil.edu", size=10.5, color=ORANGE, bold=True)
sh(s, "Text 42").top = Inches(4.82)

prs.save(OUT + ".pptx")
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", OUT + ".pptx", "--outdir", HERE],
               check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
print(OUT + ".pdf", os.path.getsize(OUT + ".pdf") // 1024, "KB")
